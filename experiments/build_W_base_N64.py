"""
build_W_base_N64.py
====================
Construye W_base N=64 desde el corpus CKM completo (Informe_Trabajo v1..v30).

Entrada (docs/):
  CKM_Informe_Trabajo.docx           (v1)
  CKM_Informe_Trabajo_v2.docx .. v28.docx
  CKM_Informe_Trabajo_v29.md
  CKM_Informe_Trabajo_v30.md

Parsing:
  docx -> doc.paragraphs, filtrando len < MIN_CHARS
  md   -> split en doble newline, excluyendo headers (#), tablas (|),
          separadores (---), filtrando len < MIN_CHARS

Nodos: NodeExtractorService(top_k=64).accumulate(paragraphs)
  -> nodes (64 terminos), pairs (co-ocurrencias), node_freq

W[i,j] = cooc[i,j] / corpus_size ; diagonal = 0 ; W simetrica

Salida: experiments/W_base_N64.json

Uso:
  python build_W_base_N64.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import numpy as np
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = ROOT / "docs"
OUT_PATH = Path(__file__).resolve().parent / "W_base_N64.json"

sys.path.insert(0, str(ROOT))
from services.node_extractor import NodeExtractorService  # noqa: E402

MIN_CHARS = 50
TOP_K = 64

MD_FILES = ["CKM_Informe_Trabajo_v29.md", "CKM_Informe_Trabajo_v30.md"]


# ---------------------------------------------------------------------------
# Descubrimiento de archivos fuente
# ---------------------------------------------------------------------------
def _version_key(path: Path) -> int:
    m = re.search(r"_v(\d+)\.docx$", path.name)
    return int(m.group(1)) if m else 1  # CKM_Informe_Trabajo.docx == v1


def discover_source_files() -> tuple[list[Path], list[Path]]:
    docx_files = sorted(
        DOCS_DIR.glob("CKM_Informe_Trabajo*.docx"), key=_version_key
    )
    md_files = [DOCS_DIR / name for name in MD_FILES]
    missing_md = [p for p in md_files if not p.exists()]
    if missing_md:
        raise FileNotFoundError(f"Faltan archivos .md esperados: {missing_md}")
    return docx_files, md_files


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def parse_docx(path: Path) -> list[str]:
    doc = Document(str(path))
    paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if len(text) >= MIN_CHARS:
            paragraphs.append(text)
    return paragraphs


_SEPARATOR_RE = re.compile(r"^[-—\s]+$")


def parse_md(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    blocks = text.split("\n\n")
    paragraphs = []
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        if block.startswith("#"):
            continue
        if "|" in block:
            continue
        if _SEPARATOR_RE.match(block):
            continue
        if len(block) < MIN_CHARS:
            continue
        paragraphs.append(block)
    return paragraphs


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    docx_files, md_files = discover_source_files()

    all_paragraphs: list[str] = []
    for path in docx_files:
        paras = parse_docx(path)
        all_paragraphs.extend(paras)
        print(f"  docx  {path.name:<40} paragraphs={len(paras)}")
    for path in md_files:
        paras = parse_md(path)
        all_paragraphs.extend(paras)
        print(f"  md    {path.name:<40} paragraphs={len(paras)}")

    corpus_size = len(all_paragraphs)
    print(f"\ncorpus_size total = {corpus_size}")

    extractor = NodeExtractorService(top_k=TOP_K)
    result = extractor.accumulate(all_paragraphs)

    nodes = result["nodes"]
    pairs = result["pairs"]
    node_freq = result["node_freq"]
    n_nodes = len(nodes)

    node_idx = {n: i for i, n in enumerate(nodes)}
    cooc = np.zeros((n_nodes, n_nodes))
    for (a, b), cnt in pairs.items():
        i, j = node_idx[a], node_idx[b]
        cooc[i, j] = cnt
        cooc[j, i] = cnt

    W = cooc / corpus_size
    np.fill_diagonal(W, 0.0)

    para_counts = {n: node_freq.get(n, 0) for n in nodes}

    output = {
        "nodes": nodes,
        "W": W.tolist(),
        "cooc": cooc.tolist(),
        "para_counts": para_counts,
        "corpus_size": corpus_size,
        "source": "CKM_Informe_Trabajo_v1-v30",
        "extraction_method": "tfidf_top64_NodeExtractorService",
        "termap_version": "v3_tfidf_auto",
    }

    OUT_PATH.write_text(
        json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    top10 = sorted(para_counts.items(), key=lambda kv: -kv[1])[:10]

    print(f"\nn_nodes = {n_nodes}")
    print(f"W shape = {W.shape}")
    print(f"corpus_size = {corpus_size}")
    print("\nTop-10 nodos por frecuencia (n_parrafos):")
    for name, freq in top10:
        print(f"  {name:<30} {freq}")

    print(f"\nGuardado: {OUT_PATH}")

    assert n_nodes == 64, f"esperado 64 nodos, obtenido {n_nodes}"
    assert W.shape == (64, 64), f"esperado shape (64,64), obtenido {W.shape}"
    assert corpus_size > 5145, f"corpus_size={corpus_size} no > 5145"


if __name__ == "__main__":
    main()
